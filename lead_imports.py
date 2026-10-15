"""Lossless source records for CSV/Excel lead imports. Never executes workbook formulas."""
import base64,csv,hashlib,io,json,re,zipfile
from datetime import date,datetime
from fastapi import HTTPException

ALIASES={
'name':['name','business name','name of the business','company','company name'],
'phone':['phone','phone number','business phone','no','number','telephone'],
'email':['email','email address'], 'category':['category','industry','trade'],
'city':['city','location','area'], 'website_url':['website url','website','site'],
'social_url':['social url','social','instagram','facebook'],
'source_url':['source url','listing url','links','link','url'], 'place_id':['place id']}
def normal(value):return re.sub(r'[^a-z0-9]+',' ',str(value or '').lower()).strip()
def text(value):
    if value is None:return ''
    if isinstance(value,(date,datetime)):return value.isoformat()
    if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value)
def parse_import(data):
    filename=data.filename;blob=None;sheets=[]
    if data.xlsx_base64:
        if not filename.lower().endswith('.xlsx'):raise HTTPException(400,'Choose an .xlsx workbook or a CSV file.')
        try:blob=base64.b64decode(data.xlsx_base64,validate=True)
        except Exception:raise HTTPException(400,'Could not read the Excel upload.') from None
        if len(blob)>5_000_000:raise HTTPException(400,'Choose an Excel file smaller than 5 MB.')
        try:
            with zipfile.ZipFile(io.BytesIO(blob)) as archive:
                if len(archive.infolist())>2000 or sum(i.file_size for i in archive.infolist())>50_000_000:raise ValueError('Workbook too large')
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(blob),data_only=False,keep_links=False)
            cached=load_workbook(io.BytesIO(blob),data_only=True,keep_links=False)
            for sheet in book:
                if sheet.max_row>5000 or sheet.max_column>100 or sheet.max_row*sheet.max_column>250000:raise ValueError('Worksheet too large')
                rows=[]
                for row in sheet:
                    cells=[]
                    for cell in row:
                        value=cell.value;formula=text(value) if cell.data_type=='f' else ''
                        if formula:value=cached[sheet.title][cell.coordinate].value
                        cells.append({'column':cell.column,'value':text(value),'formula':formula,'hyperlink':cell.hyperlink.target or cell.hyperlink.location if cell.hyperlink else '', 'format':cell.number_format,'comment':cell.comment.text if cell.comment else ''})
                    if any(c['value'] or c['formula'] or c['hyperlink'] or c['comment'] for c in cells):rows.append((row[0].row,cells))
                sheets.append((sheet.title,rows))
            book.close();cached.close()
        except ImportError:raise HTTPException(503,'Excel support is being installed. Wait for the latest Render deployment, then retry.') from None
        except Exception:raise HTTPException(400,'Could not read that workbook. Use an unencrypted .xlsx with at most 5,000 rows and 100 columns per sheet.') from None
    else:
        try:
            rows=[]
            for number,values in enumerate(csv.reader(io.StringIO(data.csv_text.lstrip('\ufeff'))),1):
                if number>5000 or len(values)>100:raise ValueError('CSV too large')
                cells=[{'column':i+1,'value':v,'formula':'','hyperlink':'','format':'','comment':''} for i,v in enumerate(values)]
                if any(values):rows.append((number,cells))
            sheets=[('CSV',rows)]
        except (csv.Error,ValueError):raise HTTPException(400,'Could not read CSV. Use at most 5,000 rows and 100 columns.') from None
    leads=[];warnings=[];section_rows=0;columns=[]
    for title,rows in sheets:
        header=None
        for index,(number,cells) in enumerate(rows[:25]):
            if any(normal(c['value']) in ALIASES['name'] for c in cells):header=index;break
        if header is None:
            if rows:warnings.append(title+': no business-name column; sheet retained only in the original workbook.')
            continue
        headers=[c['value'] or 'Column '+str(c['column']) for c in rows[header][1]];columns.append({'sheet':title,'columns':headers})
        indexes={field:next((i for i,h in enumerate(headers) if normal(h) in aliases),None) for field,aliases in ALIASES.items()}
        section=''
        for number,cells in rows[header+1:]:
            item={}
            for field,index in indexes.items():
                cell=cells[index] if index is not None and index<len(cells) else {}
                value=cell.get('value','')
                if field.endswith('_url') and cell.get('hyperlink'):value=cell['hyperlink']
                item[field]=value.strip()
            if not item['name']:
                filled=[c['value'] for c in cells if c['value']]
                if len(filled)==1:section=filled[0];section_rows+=1
                else:warnings.append(f'{title} row {number}: no business name; retained in the original file but not imported as a lead.')
                continue
            if normal(item['name']) in ALIASES['name']:continue
            original={'sheet':title,'row':number,'section':section,'cells':[{'header':headers[i] if i<len(headers) else 'Column '+str(i+1),**c} for i,c in enumerate(cells)]}
            item['_original']=original;leads.append(item)
            if len(leads)>2000:raise HTTPException(400,'Import at most 2,000 business rows at a time.')
    if not leads:raise HTTPException(400,'No lead rows found. Include a Business name or NAME OF THE BUSINESS column.')
    return {'leads':leads,'warnings':warnings,'section_rows':section_rows,'columns':columns,'blob':blob,'file_hash':hashlib.sha256(blob or data.csv_text.encode()).hexdigest()}
